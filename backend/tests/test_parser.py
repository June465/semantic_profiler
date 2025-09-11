import os
import unittest
from pathlib import Path
from docx import Document
from backend.core.resume_parser import parse_resume_file, ResumeParsingError

TEST_FILES_DIR = Path("test_parser_files")

class TestResumeParser(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        """Set up class-level resources: create test directory and dummy files."""
        TEST_FILES_DIR.mkdir(exist_ok=True)
        print(f"\nCreated test directory: {TEST_FILES_DIR}")

        doc = Document()
        doc.add_heading('Test Candidate - Software Developer', level=1)
        doc.add_paragraph('Email: test@example.com | Phone: 123-456-7890')
        doc.add_heading('Experience', level=2)
        doc.add_paragraph('Lead Engineer at Acme Corp (2020-Present)')
        doc.add_paragraph('- Developed scalable backend services.')
        doc.add_paragraph('- Mentored junior team members.')
        doc.add_heading('Skills', level=2)
        doc.add_paragraph('Python, FastAPI, SQL, Docker')
        cls.dummy_docx_path = TEST_FILES_DIR / "dummy_resume.docx"
        doc.save(cls.dummy_docx_path)
        print(f"Created dummy DOCX: {cls.dummy_docx_path}")

        cls.dummy_pdf_path = TEST_FILES_DIR / "dummy_resume.pdf"
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
            c = canvas.Canvas(str(cls.dummy_pdf_path), pagesize=letter)
            c.drawString(100, 750, "PDF Resume Example")
            c.drawString(100, 730, "Summary: Highly skilled software engineer with 5 years experience.")
            c.drawString(100, 710, "Projects: Built a data pipeline for analytics.")
            c.save()
            print(f"Created dummy PDF: {cls.dummy_pdf_path}")
            cls.has_reportlab = True
        except ImportError:
            print(f"reportlab not installed, skipping programmatic creation of {cls.dummy_pdf_path}. Please provide a real PDF for full testing.")
            cls.has_reportlab = False

    @classmethod
    def tearDownClass(cls):
        """Clean up class-level resources: remove test directory and dummy files."""
        if TEST_FILES_DIR.exists():
            for item in TEST_FILES_DIR.iterdir():
                if item.is_file():
                    item.unlink()
            TEST_FILES_DIR.rmdir()
            print(f"\nCleaned up test directory: {TEST_FILES_DIR}")

    def test_parse_docx_file(self):
        """Test parsing a valid DOCX file."""
        if not self.dummy_docx_path.exists():
            self.fail(f"Dummy DOCX file not found: {self.dummy_docx_path}")

        parsed_text = parse_resume_file(str(self.dummy_docx_path))
        self.assertIsNotNone(parsed_text)
        self.assertIn("Test Candidate", parsed_text)
        self.assertIn("Acme Corp", parsed_text)
        self.assertIn("Python, FastAPI", parsed_text)
        self.assertGreater(len(parsed_text), 100)
        self.assertFalse("\n\n" in parsed_text)
        self.assertFalse("\t" in parsed_text)

    def test_parse_pdf_file(self):
        """Test parsing a valid PDF file."""
        if not self.has_reportlab and not self.dummy_pdf_path.exists():
             self.skipTest(f"Dummy PDF creation skipped (reportlab not found) and no existing PDF at {self.dummy_pdf_path}")
        if not self.dummy_pdf_path.exists():
            self.fail(f"Dummy PDF file not found: {self.dummy_pdf_path}")

        parsed_text = parse_resume_file(str(self.dummy_pdf_path))
        self.assertIsNotNone(parsed_text)
        self.assertIn("PDF Resume Example", parsed_text)
        self.assertIn("5 years experience", parsed_text)
        self.assertGreater(len(parsed_text), 50)
        self.assertFalse("\n\n" in parsed_text)

    def test_parse_non_existent_file(self):
        """Test parsing a file that does not exist."""
        non_existent_path = TEST_FILES_DIR / "non_existent.pdf"
        with self.assertRaises(ResumeParsingError) as cm:
            parse_resume_file(str(non_existent_path))
        self.assertIn("File not found", str(cm.exception))

    def test_parse_unsupported_format(self):
        """Test parsing an unsupported file format."""
        unsupported_path = TEST_FILES_DIR / "unsupported.txt"
        unsupported_path.touch() 
        with self.assertRaises(ResumeParsingError) as cm:
            parse_resume_file(str(unsupported_path))
        self.assertIn("Unsupported file format", str(cm.exception))
        unsupported_path.unlink() 

    # def test_parse_corrupted_docx(self):
    #     corrupted_path = TEST_FILES_DIR / "corrupted.docx"
    #     # Manually create a truly corrupted file or use a known bad one
    #     with open(corrupted_path, "wb") as f:
    #         f.write(b"This is not a real docx file")
    #     with self.assertRaises(ResumeParsingError) as cm:
    #         parse_resume_file(str(corrupted_path))
    #     self.assertIn("Corrupted or invalid DOCX file", str(cm.exception))
    #     corrupted_path.unlink()


if __name__ == '__main__':
    unittest.main()