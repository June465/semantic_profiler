import os
import re
import pdfplumber
from docx import Document
from docx.opc.exceptions import OpcError
from typing import Optional

class ResumeParsingError(Exception):
    """Custom exception for resume parsing failures."""
    pass

def parse_resume_file(file_path: str) -> Optional[str]:
    extracted_text = ""
    file_extension = os.path.splitext(file_path)[1].lower()

    try:
        if file_extension == '.pdf':
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    
                    extracted_text += page.extract_text(x_tolerance=2, y_tolerance=2) or ""
        elif file_extension == '.docx':
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                extracted_text += paragraph.text + "\n"
            
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        extracted_text += cell.text + "\n"
        else:
            raise ResumeParsingError(f"Unsupported file format: {file_extension}. Only .pdf and .docx are supported.")

        
        cleaned_text = re.sub(r'\s+', ' ', extracted_text).strip()
        return cleaned_text

    except FileNotFoundError:
        raise ResumeParsingError(f"File not found: {file_path}")
    except OpcError: 
        raise ResumeParsingError(f"Corrupted or invalid DOCX file: {file_path}")
    except Exception as e:
        
        raise ResumeParsingError(f"Error parsing {file_path}: {e}")

if __name__ == '__main__':
    
    print("--- Testing parse_resume_file function ---")

    # try:
    #     c = canvas.Canvas("dummy_resume.pdf", pagesize=letter)
    #     c.drawString(100, 750, "Dummy PDF Resume")
    #     c.drawString(100, 730, "Experience: Worked at XYZ Corp for 2 years.")
    #     c.save()
    #     print("Created dummy_resume.pdf")
    # except ImportError:
    #     print("reportlab not installed, skipping dummy PDF creation. Please provide a real PDF.")

    try:
        doc = Document()
        doc.add_heading('John Doe - Software Engineer', level=1)
        doc.add_paragraph('Email: john.doe@example.com | Phone: 555-123-4567')
        doc.add_heading('Experience', level=2)
        doc.add_paragraph('Senior Developer - Tech Solutions Inc. (2020-Present)')
        doc.add_paragraph('- Led development of new features.')
        doc.add_paragraph('- Mentored junior developers.')
        doc.save("dummy_resume.docx")
        print("Created dummy_resume.docx")
    except Exception as e:
        print(f"Error creating dummy_resume.docx: {e}. Please provide a real DOCX.")


    test_files = [
        "dummy_resume.pdf",  
        "dummy_resume.docx", 
        "non_existent.pdf",
        "unsupported.txt"
    ]

    for test_file in test_files:
        print(f"\n--- Parsing: {test_file} ---")
        try:
            parsed_content = parse_resume_file(test_file)
            if parsed_content:
                print(f"Successfully parsed. First 200 chars:\n'{parsed_content[:200]}...'")
                print(f"Total length: {len(parsed_content)}")
            else:
                print("Parsed content is empty.")
        except ResumeParsingError as e:
            print(f"Parsing failed: {e}")
        except Exception as e:
            print(f"An unexpected error occurred: {e}")

    for f in ["dummy_resume.pdf", "dummy_resume.docx"]:
        if os.path.exists(f):
            os.remove(f)
            print(f"Removed {f}")